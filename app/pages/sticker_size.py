# app/pages/sticker_size.py

import openpyxl
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QTableWidgetItem, QMessageBox, QFileDialog,
)
from PySide6.QtCore import Qt

from components.search_bar import StandardSearchBar
from components.standard_page_header import StandardPageHeader
from components.standard_table import StandardTable
from components.sort_by_widget import SortByWidget
from components.generic_form_modal import GenericFormModal
from server.repositories.mstckr_repo import (
    fetch_all_mstckr,
    create_mstckr,
    update_mstckr,
    soft_delete_mstckr,
)

DPI = 96


# ── Column mapping ────────────────────────────────────────────────────────────
#
#   Tuple layout:
#       0   pk        (msstnm — also the display name)
#       1   h_in
#       2   w_in
#       3   h_px
#       4   w_px
#       5   dp_fg
#       6   ds_fg
#       7   pt_fg
#       8   pt_ct
#       9   pt_id
#       10  pt_dt
#       11  source
#       12  user_remark
#       13  item_remark
#       14  added_by
#       15  added_at
#       16  changed_by
#       17  changed_at
#       18  changed_no

_COL_HEADER_TO_TUPLE_IDX = {
    "NAME":           0,
    "WIDTH (INCH)":   2,
    "WIDTH (PX)":     4,
    "HEIGHT (INCH)":  1,
    "HEIGHT (PX)":    3,
    "ADDED BY":       14,
    "ADDED AT":       15,
    "CHANGED BY":     16,
    "CHANGED AT":     17,
    "CHANGED NO":     18,
}

_TABLE_HEADERS = [
    "NAME",
    "WIDTH (INCH)", "WIDTH (PX)",
    "HEIGHT (INCH)", "HEIGHT (PX)",
    "ADDED BY", "ADDED AT", "CHANGED BY", "CHANGED AT", "CHANGED NO",
]


# ── Form schema ───────────────────────────────────────────────────────────────

def _build_form_schema() -> list[dict]:
    return [
        {
            "name":        "name",
            "label":       "Sticker Name",
            "type":        "text",
            "placeholder": "Enter sticker name (max 20 chars)",
            "required":    True,
        },
        {
            "name":     "width",
            "label":    "Width",
            "type":     "dimension_pair",
            "dpi":      DPI,
            "required": True,
        },
        {
            "name":     "height",
            "label":    "Height",
            "type":     "dimension_pair",
            "dpi":      DPI,
            "required": True,
        },
        # ── Audit (always readonly) ────────────────────────────────────
        {"name": "added_by",   "label": "Added By",   "type": "readonly"},
        {"name": "added_at",   "label": "Added At",   "type": "readonly"},
        {"name": "changed_by", "label": "Changed By", "type": "readonly"},
        {"name": "changed_at", "label": "Changed At", "type": "readonly"},
        {"name": "changed_no", "label": "Changed No", "type": "readonly"},
    ]


# ── Data conversion ───────────────────────────────────────────────────────────

def _row_to_tuple(r: dict) -> tuple:
    return (
        (r.get("pk") or "").strip(),                                    # 0  name/pk
        float(r.get("h_in") or 0),                                      # 1  h_in
        float(r.get("w_in") or 0),                                      # 2  w_in
        int(r.get("h_px") or 0),                                        # 3  h_px
        int(r.get("w_px") or 0),                                        # 4  w_px
        (r.get("dp_fg") or "").strip(),                                 # 5  dp_fg
        (r.get("ds_fg") or "").strip(),                                 # 6  ds_fg
        (r.get("pt_fg") or "").strip(),                                 # 7  pt_fg
        int(r.get("pt_ct") or 0),                                       # 8  pt_ct
        (r.get("pt_id") or "").strip(),                                 # 9  pt_id
        str(r["pt_dt"])[:19] if r.get("pt_dt") else "",                 # 10 pt_dt
        (r.get("source") or "").strip(),                                # 11 source
        (r.get("user_remark") or "").strip(),                           # 12 user_remark
        (r.get("item_remark") or "").strip(),                           # 13 item_remark
        (r.get("added_by") or "").strip(),                              # 14 added_by
        str(r["added_at"])[:19] if r.get("added_at") else "",           # 15 added_at
        (r.get("changed_by") or "").strip(),                            # 16 changed_by
        str(r["changed_at"])[:19] if r.get("changed_at") else "",       # 17 changed_at
        str(r.get("changed_no", 0)),                                    # 18 changed_no
    )


def _parse_submitted_dims(data: dict) -> tuple[float, float, int, int] | None:
    try:
        w_in = float(data["width_in"])
        w_px = int(data["width_px"])
        h_in = float(data["height_in"])
        h_px = int(data["height_px"])
    except (KeyError, ValueError, TypeError):
        return None
    if w_in <= 0 or w_px <= 0 or h_in <= 0 or h_px <= 0:
        return None
    return w_in, h_in, w_px, h_px


# ── Page ──────────────────────────────────────────────────────────────────────

class StickerSizePage(QWidget):
    def __init__(self):
        super().__init__()
        self.all_data: list[tuple] = []
        self.filtered_data: list[tuple] = []
        self.current_page = 0
        self.page_size = 25
        self.available_page_sizes = [25, 50, 100]
        self._last_filter_type = "NAME"
        self._last_search_text = ""
        self._sort_fields: list[str] = []
        self._sort_directions: dict[str, str] = {}
        self._active_modal: GenericFormModal | None = None
        self._init_ui()
        self.load_data()

    # ── UI setup ──────────────────────────────────────────────────────────────

    def _init_ui(self):
        self.setStyleSheet("background-color: #F8FAFC;")
        layout = QVBoxLayout(self)
        layout.setContentsMargins(40, 20, 40, 12)
        layout.setSpacing(0)
        self.main_layout = layout

        self.header = StandardPageHeader(
            title="Sticker Size",
            subtitle="Manage sticker size definitions (width x height).",
            enabled_actions=["Add", "Excel", "Refresh", "View Detail"],
        )
        layout.addWidget(self.header)
        self._connect_header_actions()
        layout.addSpacing(12)

        self.search_bar = StandardSearchBar()
        self.search_bar.searchChanged.connect(self.filter_table)
        layout.addWidget(self.search_bar)
        layout.addSpacing(5)

        self.table_comp = StandardTable(_TABLE_HEADERS)
        self.table = self.table_comp.table

        self.sort_bar = SortByWidget(self.table)
        self.sort_bar.sortChanged.connect(self.on_sort_changed)
        layout.addWidget(self.sort_bar)
        layout.addSpacing(8)

        layout.addWidget(self.table_comp)
        layout.addSpacing(16)

        self.pagination = self.table_comp.pagination
        self.pagination.pageChanged.connect(self.on_page_changed)
        self.pagination.pageSizeChanged.connect(self.on_page_size_changed)

        self.sort_bar.initialize_default_sort()
        self.table.itemSelectionChanged.connect(self._on_row_selection_changed)
        self._update_selection_dependent_state(False)

    # ── Selection helpers ─────────────────────────────────────────────────────

    def _on_row_selection_changed(self):
        self._update_selection_dependent_state(bool(self.table.selectedItems()))

    def _update_selection_dependent_state(self, enabled: bool):
        if self._active_modal is not None:
            return
        for label in ("Edit", "Delete", "View Detail"):
            btn = self.header.get_action_button(label)
            if btn:
                btn.setEnabled(enabled)

    def _get_selected_row(self) -> tuple | None:
        rows = self.table.selectionModel().selectedRows()
        if not rows:
            return None
        global_idx = self.current_page * self.page_size + rows[0].row()
        if global_idx >= len(self.filtered_data):
            return None
        return self.filtered_data[global_idx]

    # ── Modal helpers ─────────────────────────────────────────────────────────

    _ALL_HEADER_ACTIONS = ["Add", "Excel", "Refresh", "Edit", "Delete", "View Detail"]

    def _lock_header(self):
        for label in self._ALL_HEADER_ACTIONS:
            btn = self.header.get_action_button(label)
            if btn:
                btn.setEnabled(False)

    def _unlock_header(self):
        has_sel = bool(self.table.selectedItems())
        for label in self._ALL_HEADER_ACTIONS:
            btn = self.header.get_action_button(label)
            if btn:
                btn.setEnabled(has_sel if label in ("Edit", "Delete", "View Detail") else True)

    def _clear_active_modal(self):
        self._active_modal = None

    def _open_modal(self, modal: GenericFormModal):
        modal.opened.connect(self._lock_header)
        modal.closed.connect(self._unlock_header)
        modal.closed.connect(self._clear_active_modal)
        self._active_modal = modal
        modal.exec()

    # ── Page visibility ───────────────────────────────────────────────────────

    def showEvent(self, event):
        super().showEvent(event)
        if self._active_modal is not None and not self._active_modal.isVisible():
            self._active_modal.show()

    def hideEvent(self, event):
        super().hideEvent(event)
        if self._active_modal is not None and self._active_modal.isVisible():
            self._active_modal.hide()

    # ── Rendering ─────────────────────────────────────────────────────────────

    def _make_item(self, text: str) -> QTableWidgetItem:
        it = QTableWidgetItem(text)
        it.setTextAlignment(Qt.AlignTop | Qt.AlignLeft)
        return it

    def _add_table_row(self, row: tuple):
        r = self.table.rowCount()
        self.table.insertRow(r)
        self.table.setItem(r, 0, self._make_item(row[0]))           # name/pk
        self.table.setItem(r, 1, self._make_item(str(row[2])))      # w_in
        self.table.setItem(r, 2, self._make_item(str(row[4])))      # w_px
        self.table.setItem(r, 3, self._make_item(str(row[1])))      # h_in
        self.table.setItem(r, 4, self._make_item(str(row[3])))      # h_px
        self.table.setItem(r, 5, self._make_item(row[14]))          # added_by
        self.table.setItem(r, 6, self._make_item(row[15]))          # added_at
        self.table.setItem(r, 7, self._make_item(row[16]))          # changed_by
        self.table.setItem(r, 8, self._make_item(row[17]))          # changed_at
        self.table.setItem(r, 9, self._make_item(row[18]))          # changed_no

    def render_page(self):
        self.table.setSortingEnabled(False)
        self.table.setRowCount(0)

        data = self.filtered_data or []
        total = len(data)
        start = self.current_page * self.page_size
        end = min(start + self.page_size, total)

        for item in data[start:end]:
            self._add_table_row(item)

        for r in range(end - start):
            self.table.setVerticalHeaderItem(r, QTableWidgetItem(str(start + r + 1)))

        self.table.resizeRowsToContents()
        self.pagination.update(
            start=0 if total == 0 else start + 1,
            end=0 if total == 0 else end,
            total=total,
            has_prev=self.current_page > 0,
            has_next=end < total,
            current_page=self.current_page,
            page_size=self.page_size,
            available_page_sizes=self.available_page_sizes,
        )

    # ── Data loading ──────────────────────────────────────────────────────────

    def load_data(self):
        try:
            self.all_data = [_row_to_tuple(r) for r in fetch_all_mstckr()]
        except Exception as exc:
            QMessageBox.critical(self, "Database Error", f"Failed to load data:\n\n{exc}")
            self.all_data = []
        self._apply_filter_and_reset_page()

    # ── Filtering & sorting ───────────────────────────────────────────────────

    def filter_table(self, filter_type: str, search_text: str):
        self._last_filter_type = filter_type
        self._last_search_text = search_text
        self._apply_filter_and_reset_page()

    def _apply_filter_and_reset_page(self):
        query = (self._last_search_text or "").lower().strip()
        col_idx = _COL_HEADER_TO_TUPLE_IDX.get(self._last_filter_type, 0)
        self.filtered_data = (
            list(self.all_data) if not query
            else [row for row in self.all_data if query in str(row[col_idx] or "").lower()]
        )
        self._apply_sort()
        self.current_page = 0
        self.render_page()

    def on_sort_changed(self, fields: list[str], field_directions: dict):
        self._sort_fields = fields or []
        self._sort_directions = field_directions or {}
        self._apply_filter_and_reset_page()

    def _apply_sort(self):
        if not self._sort_fields or not self.filtered_data:
            return
        for field in reversed(self._sort_fields):
            idx = _COL_HEADER_TO_TUPLE_IDX.get(field)
            if idx is None:
                continue
            self.filtered_data.sort(
                key=lambda row, i=idx: self._sort_key(row, i),
                reverse=(self._sort_directions.get(field, "asc") == "desc"),
            )

    def _sort_key(self, row: tuple, idx: int):
        val = str(row[idx]) if row[idx] is not None else ""
        try:
            return float(val.replace(",", ""))
        except ValueError:
            return val.lower()

    # ── Pagination ────────────────────────────────────────────────────────────

    def on_page_changed(self, page_action: int):
        total = len(self.filtered_data or [])
        total_pages = max(1, (total + self.page_size - 1) // self.page_size)
        if page_action == -1:
            self.current_page = max(0, self.current_page - 1)
        elif page_action == 1:
            self.current_page = min(total_pages - 1, self.current_page + 1)
        else:
            self.current_page = max(0, min(int(page_action), total_pages - 1))
        self.render_page()

    def on_page_size_changed(self, new_size: int):
        self.page_size = new_size
        self.current_page = 0
        self.render_page()

    # ── Header button wiring ──────────────────────────────────────────────────

    def _connect_header_actions(self):
        for label, slot in {
            "Refresh":     self.load_data,
            "Add":         self.handle_add_action,
            "Excel":       self.handle_export_action,
            "Edit":        self.handle_edit_action,
            "Delete":      self.handle_delete_action,
            "View Detail": self.handle_view_detail_action,
        }.items():
            btn = self.header.get_action_button(label)
            if btn:
                btn.clicked.connect(slot)

    # ── Shared initial_data builder ───────────────────────────────────────────

    def _row_to_modal_data(self, row: tuple) -> dict:
        return {
            "name":       row[0],
            "width_in":   str(row[2]),
            "width_px":   str(row[4]),
            "height_in":  str(row[1]),
            "height_px":  str(row[3]),
            "added_by":   row[14],
            "added_at":   row[15],
            "changed_by": row[16],
            "changed_at": row[17],
            "changed_no": row[18],
        }

    # ── Action handlers ───────────────────────────────────────────────────────

    def handle_add_action(self):
        modal = GenericFormModal(
            title="Add Sticker Size",
            fields=_build_form_schema(),
            parent=self,
            mode="add",
        )
        modal.formSubmitted.connect(self._on_add_submitted)
        self._open_modal(modal)

    def _on_add_submitted(self, data: dict):
        name = data.get("name", "").strip()
        if not name:
            QMessageBox.warning(self, "Validation", "Sticker name is required.")
            return
        if len(name) > 20:
            QMessageBox.warning(self, "Validation", "Sticker name must be 20 characters or fewer.")
            return
        dims = _parse_submitted_dims(data)
        if dims is None:
            QMessageBox.warning(self, "Validation",
                                "Width and height must be valid positive numbers.")
            return
        w_in, h_in, w_px, h_px = dims
        try:
            create_mstckr(name=name, h_in=h_in, w_in=w_in, h_px=h_px, w_px=w_px)
        except Exception as exc:
            QMessageBox.critical(self, "Database Error", f"Insert failed:\n\n{exc}")
            return
        self.load_data()

    def handle_edit_action(self):
        row = self._get_selected_row()
        if row is None:
            return
        modal = GenericFormModal(
            title="Edit Sticker Size",
            fields=_build_form_schema(),
            parent=self,
            mode="edit",
            initial_data=self._row_to_modal_data(row),
        )
        modal.formSubmitted.connect(lambda data, r=row: self._on_edit_submitted(r, data))
        self._open_modal(modal)

    def _on_edit_submitted(self, row: tuple, data: dict):
        dims = _parse_submitted_dims(data)
        if dims is None:
            QMessageBox.warning(self, "Validation",
                                "Width and height must be valid positive numbers.")
            return
        w_in, h_in, w_px, h_px = dims
        old_changed_no = int(row[18]) if str(row[18]).isdigit() else 0
        try:
            new_name = data.get("name", "").strip()

            update_mstckr(
                old_pk=row[0],
                new_name=new_name,
                h_in=h_in,
                w_in=w_in,
                h_px=h_px,
                w_px=w_px,
                old_changed_no=old_changed_no,
            )
        except Exception as exc:
            QMessageBox.critical(self, "Database Error", f"Update failed:\n\n{exc}")
            return
        self.load_data()

    def handle_view_detail_action(self):
        row = self._get_selected_row()
        if row is None:
            return
        modal = GenericFormModal(
            title="Sticker Size Detail",
            subtitle="Full details for the selected record.",
            fields=_build_form_schema(),
            parent=self,
            mode="view",
            initial_data=self._row_to_modal_data(row),
        )
        self._open_modal(modal)

    def handle_export_action(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Save Excel File", "sticker_size.xlsx", "Excel Files (*.xlsx)"
        )
        if not path:
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sticker Size"
        ws.append(_TABLE_HEADERS)
        for row in self.filtered_data:
            ws.append([
                row[0],                           # name
                row[2],  row[4],                  # w_in, w_px
                row[1],  row[3],                  # h_in, h_px
                row[14], row[15],                 # added_by, added_at
                row[16], row[17], row[18],         # changed_by, changed_at, changed_no
            ])
        wb.save(path)
        QMessageBox.information(self, "Export Complete",
                                f"Exported {len(self.filtered_data)} records to:\n{path}")

    def handle_delete_action(self):
        row = self._get_selected_row()
        if row is None:
            return
        msg = QMessageBox(self)
        msg.setWindowTitle("Confirm Delete")
        msg.setText("Are you sure you want to delete this record?")
        msg.setInformativeText(f"Name: {row[0]}")
        msg.setStandardButtons(QMessageBox.Yes | QMessageBox.Cancel)
        msg.setDefaultButton(QMessageBox.Cancel)
        msg.setIcon(QMessageBox.Warning)
        if msg.exec() == QMessageBox.Yes:
            try:
                soft_delete_mstckr(row[0])
            except Exception as exc:
                QMessageBox.critical(self, "Database Error", f"Delete failed:\n\n{exc}")
                return
            self.load_data()